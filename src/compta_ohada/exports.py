"""
Export des états en Excel (.xlsx) et en PDF.

    compta.definir_parametre("raison_sociale", "Quincaillerie Alpha SARL")
    compta.definir_parametre("nif", "084123456M")
    compta.exporter_etat("bilan", "bilan_2026.pdf", exercice="2026")
    compta.exporter_etat("balance", "balance.xlsx", debut="2026-01-01", fin="2026-12-31")

Dépendances optionnelles : ``pip install compta-ohada[exports]`` (openpyxl,
reportlab) — les mêmes que les ERP Beyrey embarquent déjà.

Chaque état est d'abord mis sous une forme tabulaire neutre (``Document`` :
blocs, colonnes typées, lignes avec style), puis rendu dans le format voulu.
Les montants restent NUMÉRIQUES dans Excel (formules et tris possibles).
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from .erreurs import ParametreInvalide
from .socle import ZERO

TEXTE, MONTANT, DATE, POURCENT = "texte", "montant", "date", "pourcent"


@dataclass
class Bloc:
    titre: str | None
    colonnes: list[tuple[str, str]]                 # (intitulé, type)
    lignes: list[tuple[list[Any], str | None]] = field(default_factory=list)
    # style : None, "total", "sous_total", "titre"

    def ajouter(self, valeurs: list[Any], style: str | None = None) -> None:
        self.lignes.append((valeurs, style))


@dataclass
class Document:
    titre: str
    sous_titre: str
    blocs: list[Bloc]
    paysage: bool = False


# ---------------------------------------------------------------------------
# États -> Document
# ---------------------------------------------------------------------------

def _d(valeur) -> str:
    if isinstance(valeur, (date, datetime)):
        return valeur.strftime("%d/%m/%Y")
    return str(valeur or "")


def doc_balance(e: dict) -> Document:
    b = Bloc(None, [("Compte", TEXTE), ("Libellé", TEXTE),
                    ("Ouverture débit", MONTANT), ("Ouverture crédit", MONTANT),
                    ("Mouvements débit", MONTANT), ("Mouvements crédit", MONTANT),
                    ("Solde débiteur", MONTANT), ("Solde créditeur", MONTANT)])
    for l in e["lignes"]:
        b.ajouter([l["compte"], l["libelle"], l["ouverture_debit"],
                   l["ouverture_credit"], l["debit"], l["credit"],
                   l["solde_debit"], l["solde_credit"]],
                  "sous_total" if len(l["compte"]) <= 2 else None)
    t = e["totaux"]
    b.ajouter(["", "TOTAUX", t["ouverture_debit"], t["ouverture_credit"],
               t["debit"], t["credit"], t["solde_debit"], t["solde_credit"]], "total")
    return Document("Balance générale",
                    f"Exercice {e['exercice']} — du {_d(e['debut'])} au {_d(e['fin'])}",
                    [b], paysage=True)


def doc_grand_livre(livre: list[dict], debut, fin) -> Document:
    b = Bloc(None, [("Date", DATE), ("Journal", TEXTE), ("Pièce", TEXTE),
                    ("Libellé", TEXTE), ("Débit", MONTANT), ("Crédit", MONTANT),
                    ("Solde", MONTANT)])
    for c in livre:
        b.ajouter([None, "", c["compte"], c["libelle"], None, None, None], "titre")
        b.ajouter([None, "", "", "Solde d'ouverture", None, None, c["ouverture"]])
        for m in c["mouvements"]:
            b.ajouter([m["date"], m["journal"], m["numero"], m["libelle"],
                       m["debit"], m["credit"], m["solde"]])
        b.ajouter([None, "", "", f"Total {c['compte']}", c["total_debit"],
                   c["total_credit"], c["solde"]], "sous_total")
    return Document("Grand livre", f"Du {_d(debut)} au {_d(fin)}", [b], paysage=True)


def doc_journal(j: dict) -> Document:
    b = Bloc(None, [("Date", DATE), ("Pièce", TEXTE), ("Compte", TEXTE),
                    ("Libellé", TEXTE), ("Tiers", TEXTE), ("Débit", MONTANT),
                    ("Crédit", MONTANT)])
    for e in j["ecritures"]:
        for i, l in enumerate(e["lignes"]):
            b.ajouter([e["date"] if i == 0 else None, e["numero"] if i == 0 else "",
                       l["compte"], l["libelle"], l["tiers_ref"] or "",
                       l["debit"] or None, l["credit"] or None])
    b.ajouter([None, "", "", "TOTAUX", "", j["total_debit"], j["total_credit"]], "total")
    return Document(f"Journal {j['journal']}",
                    f"Du {_d(j['debut'])} au {_d(j['fin'])}", [b], paysage=True)


def doc_balance_tiers(lignes: list[dict], prefixe: str, fin) -> Document:
    b = Bloc(None, [("Tiers", TEXTE), ("Nom", TEXTE), ("Débit", MONTANT),
                    ("Crédit", MONTANT), ("Solde", MONTANT)])
    total = ZERO
    for l in lignes:
        b.ajouter([l["tiers_ref"] or "", l["tiers_nom"] or "", l["debit"],
                   l["credit"], l["solde"]])
        total += l["solde"]
    b.ajouter(["", "TOTAL", None, None, total], "total")
    return Document(f"Balance des tiers — comptes {prefixe}", f"Au {_d(fin)}", [b])


def doc_balance_agee(lignes: list[dict], prefixe: str, jour) -> Document:
    tranches = [k for k in (lignes[0].keys() if lignes else [])
                if k == "non_echu" or "_" in k and k[0].isdigit() or k.startswith("plus_")]
    if not tranches:
        tranches = ["non_echu", "0_30", "31_60", "61_90", "plus_90"]
    noms = {"non_echu": "Non échu"}
    for t in tranches:
        if t.startswith("plus_"):
            noms[t] = f"+ {t[5:]} j"
        elif t != "non_echu":
            noms[t] = t.replace("_", " à ") + " j"
    b = Bloc(None, [("Compte", TEXTE), ("Tiers", TEXTE), ("Nom", TEXTE)]
             + [(noms[t], MONTANT) for t in tranches] + [("Solde", MONTANT)])
    totaux = [ZERO] * (len(tranches) + 1)
    for l in lignes:
        valeurs = [l[t] for t in tranches] + [l["solde"]]
        totaux = [a + v for a, v in zip(totaux, valeurs)]
        b.ajouter([l["compte"], l["tiers_ref"] or "", l["tiers_nom"] or ""] + valeurs)
    b.ajouter(["", "", "TOTAL"] + totaux, "total")
    return Document(f"Balance âgée — comptes {prefixe}", f"Au {_d(jour)}", [b],
                    paysage=True)


def doc_declaration_tva(d: dict) -> Document:
    synthese = Bloc("Synthèse", [("Rubrique", TEXTE), ("Montant", MONTANT)])
    for lib, cle in (("TVA collectée", "tva_collectee"),
                     ("TVA déductible", "tva_deductible"),
                     ("Crédit de TVA antérieur", "credit_anterieur")):
        synthese.ajouter([lib, d[cle]])
    synthese.ajouter(["TVA NETTE À PAYER", d["tva_due"]], "total")
    synthese.ajouter(["CRÉDIT DE TVA À REPORTER", d["credit_a_reporter"]], "total")
    detail = Bloc("Détail par taux", [("Code", TEXTE), ("Libellé", TEXTE),
                                       ("Taux %", POURCENT), ("Base HT", MONTANT),
                                       ("TVA", MONTANT)])
    for x in d["detail"]:
        detail.ajouter([x["code"], x["libelle"], x["taux"], x["base"], x["tva"]])
    return Document("Déclaration de TVA",
                    f"Du {_d(d['debut'])} au {_d(d['fin'])}", [synthese, detail])


def doc_bilan(bi: dict) -> Document:
    actif = Bloc("ACTIF", [("Réf.", TEXTE), ("Actif", TEXTE), ("Brut", MONTANT),
                           ("Amort./Dépréc.", MONTANT), ("Net", MONTANT),
                           ("Net N-1", MONTANT)])
    for l in bi["actif"]:
        actif.ajouter([l["code"], l["libelle"], l["brut"], l["amort"], l["net"],
                       l["net_n1"]], "total" if l["total"] else None)
    passif = Bloc("PASSIF", [("Réf.", TEXTE), ("Passif", TEXTE), ("Net", MONTANT),
                             ("Net N-1", MONTANT)])
    for l in bi["passif"]:
        passif.ajouter([l["code"], l["libelle"], l["net"], l["net_n1"]],
                       "total" if l["total"] else None)
    return Document("Bilan", f"Exercice {bi['exercice']} — arrêté au "
                    f"{_d(bi['date_arrete'])}", [actif, passif])


def doc_compte_de_resultat(cr: dict) -> Document:
    b = Bloc(None, [("Réf.", TEXTE), ("Libellé", TEXTE), ("Exercice", MONTANT),
                    ("Exercice N-1", MONTANT)])
    for l in cr["lignes"]:
        b.ajouter([l["code"], l["libelle"], l["montant"], l["montant_n1"]],
                  "total" if l["solde_intermediaire"] else None)
    return Document("Compte de résultat",
                    f"Exercice {cr['exercice']} — du {_d(cr['debut'])} au {_d(cr['fin'])}",
                    [b])


def doc_tft(t: dict) -> Document:
    b = Bloc(None, [("Réf.", TEXTE), ("Libellé", TEXTE), ("Montant", MONTANT)])
    for l in t["lignes"]:
        b.ajouter([l["code"], l["libelle"], l["montant"]],
                  "total" if l["total"] else None)
    return Document("Tableau des flux de trésorerie", f"Exercice {t['exercice']}", [b])


def doc_tableau_immobilisations(t: dict) -> Document:
    cles = ("brut_debut", "entrees", "sorties", "brut_fin", "amort_debut",
            "dotations", "reprises", "amort_fin", "vnc")
    b = Bloc(None, [("Code", TEXTE), ("Libellé", TEXTE), ("Compte", TEXTE),
                    ("Brut début", MONTANT), ("Entrées", MONTANT), ("Sorties", MONTANT),
                    ("Brut fin", MONTANT), ("Amort. début", MONTANT),
                    ("Dotations", MONTANT), ("Reprises", MONTANT),
                    ("Amort. fin", MONTANT), ("VNC", MONTANT)])
    for l in t["lignes"]:
        b.ajouter([l["code"], l["libelle"], l["compte"]] + [l[k] for k in cles])
    b.ajouter(["", "TOTAUX", ""] + [t["totaux"][k] for k in cles], "total")
    return Document("Tableau des immobilisations et des amortissements",
                    f"Exercice {t['exercice']}", [b], paysage=True)


def doc_plan_amortissement(plan: list[dict], code: str) -> Document:
    b = Bloc(None, [("Exercice", TEXTE), ("Du", DATE), ("Au", DATE),
                    ("Base", MONTANT), ("Dotation", MONTANT), ("Cumul", MONTANT),
                    ("VNC", MONTANT)])
    for a in plan:
        b.ajouter([str(a["exercice"]), a["debut"], a["fin"], a["base"],
                   a["dotation"], a["cumul"], a["vnc"]])
    return Document("Plan d'amortissement", f"Immobilisation {code}", [b])


def doc_etats_smt(e: dict) -> Document:
    bilan = Bloc("Bilan (système minimal de trésorerie)",
                 [("Réf.", TEXTE), ("Poste", TEXTE), ("Montant", MONTANT)])
    for l in e["bilan"]:
        bilan.ajouter([l["code"], l["libelle"], l["montant"]],
                      "total" if l["code"] in ("AT", "PT") else None)
    cr = Bloc("Compte de résultat (système minimal de trésorerie)",
              [("Réf.", TEXTE), ("Poste", TEXTE), ("Montant", MONTANT)])
    for l in e["compte_de_resultat"]:
        cr.ajouter([l["code"], l["libelle"], l["montant"]],
                   "total" if l["code"] in ("RT", "DT", "RN") else None)
    return Document("États financiers — Système minimal de trésorerie",
                    f"Exercice {e['exercice']}", [bilan, cr])


def doc_livre_recettes_depenses(l: dict) -> Document:
    b = Bloc(None, [("Date", DATE), ("Pièce", TEXTE), ("Libellé", TEXTE),
                    ("Nature", TEXTE), ("Recette", MONTANT), ("Dépense", MONTANT),
                    ("Solde", MONTANT)])
    b.ajouter([None, "", "Solde initial", "", None, None, l["solde_initial"]],
              "sous_total")
    for m in l["mouvements"]:
        b.ajouter([m["date"], m["numero"], m["libelle"], m["nature"],
                   m["recette"] or None, m["depense"] or None, m["solde"]])
    b.ajouter([None, "", "TOTAUX", "", l["total_recettes"], l["total_depenses"],
               l["solde_final"]], "total")
    return Document("Livre des recettes et des dépenses",
                    f"Du {_d(l['debut'])} au {_d(l['fin'])}", [b], paysage=True)


def doc_resultat_analytique(r: dict) -> Document:
    b = Bloc(None, [("Section", TEXTE), ("Libellé", TEXTE), ("Produits", MONTANT),
                    ("Charges", MONTANT), ("Résultat", MONTANT), ("Marge %", POURCENT)])
    for s in r["sections"]:
        b.ajouter([s["section"], s["libelle"], s["produits"], s["charges"],
                   s["resultat"], s["marge_pct"]])
    t = r["total"]
    b.ajouter(["", "TOTAL", t["produits"], t["charges"], t["resultat"], None], "total")
    return Document("Résultat analytique" + (f" — axe {r['axe']}" if r["axe"] else ""),
                    f"Du {_d(r['debut'])} au {_d(r['fin'])}", [b])


# ---------------------------------------------------------------------------
# Rendus
# ---------------------------------------------------------------------------

def _format_montant(v: Decimal, decimales: int) -> str:
    if v is None:
        return ""
    v = Decimal(v)
    texte = f"{abs(v):,.{decimales}f}".replace(",", " ").replace(".", ",")
    return f"-{texte}" if v < 0 else texte


def rendre_xlsx(doc: Document, chemin: Path, entete: list[str], decimales: int) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover
        raise ParametreInvalide(
            "Export Excel : installez openpyxl (pip install compta-ohada[exports]).") from exc
    wb = Workbook()
    ws = wb.active
    ws.title = doc.titre[:31].replace("/", "-")
    gras = Font(bold=True)
    fond_total = PatternFill("solid", fgColor="E8EEF4")
    fond_entete = PatternFill("solid", fgColor="1F3A5F")
    blanc = Font(bold=True, color="FFFFFF")
    fin = Side(style="thin", color="B7C3D0")
    format_nombre = "#,##0" if decimales == 0 else "#,##0.00"
    rang = 1
    for ligne_entete in entete:
        ws.cell(rang, 1, ligne_entete).font = Font(bold=rang == 1, size=11)
        rang += 1
    ws.cell(rang, 1, doc.titre).font = Font(bold=True, size=14)
    rang += 1
    ws.cell(rang, 1, doc.sous_titre).font = Font(italic=True)
    rang += 2
    largeurs: dict[int, int] = {}
    premier_tableau = None
    for bloc in doc.blocs:
        if bloc.titre:
            ws.cell(rang, 1, bloc.titre).font = Font(bold=True, size=12)
            rang += 1
        for j, (intitule, _) in enumerate(bloc.colonnes, 1):
            c = ws.cell(rang, j, intitule)
            c.font, c.fill = blanc, fond_entete
            c.alignment = Alignment(horizontal="center", vertical="center",
                                    wrap_text=True)
            largeurs[j] = max(largeurs.get(j, 8), len(intitule) + 2)
        if premier_tableau is None:
            premier_tableau = rang
        rang += 1
        for valeurs, style in bloc.lignes:
            for j, ((_, type_), v) in enumerate(zip(bloc.colonnes, valeurs), 1):
                if v is None or v == "":
                    cell = ws.cell(rang, j)
                elif type_ == MONTANT:
                    cell = ws.cell(rang, j, float(Decimal(v)))
                    cell.number_format = format_nombre
                elif type_ == POURCENT:
                    cell = ws.cell(rang, j, float(Decimal(v)))
                    cell.number_format = "0.0"
                elif type_ == DATE and isinstance(v, (date, datetime)):
                    cell = ws.cell(rang, j, v)
                    cell.number_format = "DD/MM/YYYY"
                else:
                    cell = ws.cell(rang, j, str(v))
                cell.border = Border(bottom=fin)
                if style in ("total", "sous_total", "titre"):
                    cell.font = gras
                if style == "total":
                    cell.fill = fond_total
                if v not in (None, ""):
                    longueur = len(_format_montant(v, decimales)) if type_ == MONTANT \
                        else len(_d(v) if type_ == DATE else str(v))
                    largeurs[j] = min(max(largeurs.get(j, 8), longueur + 2), 60)
            rang += 1
        rang += 1
    for j, l in largeurs.items():
        ws.column_dimensions[get_column_letter(j)].width = l
    if premier_tableau:
        ws.freeze_panes = ws.cell(premier_tableau + 1, 1)
    ws.page_setup.orientation = "landscape" if doc.paysage else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
    wb.properties.creator = "compta-ohada"
    wb.properties.title = doc.titre
    wb.save(chemin)


def _pdf_texte(texte: str) -> str:
    """Les polices PDF standard couvrent Windows-1252 : on remplace le reste."""
    texte = str(texte).replace(" ", " ").replace(" ", " ")
    return texte.encode("cp1252", "replace").decode("cp1252")


def rendre_pdf(doc: Document, chemin: Path, entete: list[str], decimales: int) -> None:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import (LongTable, Paragraph, SimpleDocTemplate,
                                        Spacer, TableStyle)
    except ImportError as exc:  # pragma: no cover
        raise ParametreInvalide(
            "Export PDF : installez reportlab (pip install compta-ohada[exports]).") from exc
    taille = landscape(A4) if doc.paysage else A4
    styles = getSampleStyleSheet()
    cellule = ParagraphStyle("cellule", parent=styles["BodyText"], fontSize=7.5,
                             leading=9)
    cellule_gras = ParagraphStyle("cellule_gras", parent=cellule,
                                  fontName="Helvetica-Bold")
    entete_col = ParagraphStyle("entete_col", parent=cellule_gras,
                                textColor=colors.white)
    entete_col_droite = ParagraphStyle("entete_col_droite", parent=entete_col,
                                       alignment=2)
    largeur_utile = taille[0] - 24 * mm
    histoire = []
    for i, l in enumerate(entete):
        histoire.append(Paragraph(_pdf_texte(l), styles["Heading4" if i == 0 else "Normal"]))
    histoire.append(Spacer(1, 3 * mm))
    histoire.append(Paragraph(_pdf_texte(doc.titre), styles["Title"]))
    histoire.append(Paragraph(_pdf_texte(doc.sous_titre), styles["Italic"]))
    histoire.append(Spacer(1, 4 * mm))
    for bloc in doc.blocs:
        if bloc.titre:
            histoire.append(Paragraph(_pdf_texte(bloc.titre), styles["Heading3"]))
        poids = [3.2 if t == TEXTE and i == 1 or (t == TEXTE and "Libell" in n)
                 else 1.0 if t in (DATE, POURCENT) else 1.4
                 for i, (n, t) in enumerate(bloc.colonnes)]
        total_poids = sum(poids)
        largeurs = [largeur_utile * p / total_poids for p in poids]
        donnees = [[Paragraph(_pdf_texte(n), entete_col_droite
                              if t in (MONTANT, POURCENT) else entete_col)
                    for n, t in bloc.colonnes]]
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 7.5),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#B7C3D0")),
            ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]
        for r, (valeurs, st) in enumerate(bloc.lignes, 1):
            ligne = []
            for (n, t), v in zip(bloc.colonnes, valeurs):
                if t == MONTANT:
                    texte = _format_montant(v, decimales) if v is not None else ""
                elif t == POURCENT:
                    texte = "" if v is None else f"{Decimal(v):.1f}".replace(".", ",")
                else:
                    texte = _d(v) if t == DATE else str(v or "")
                ligne.append(Paragraph(_pdf_texte(texte),
                                       cellule_gras if st else cellule)
                             if t == TEXTE else _pdf_texte(texte))
            donnees.append(ligne)
            if st == "total":
                style.append(("BACKGROUND", (0, r), (-1, r), colors.HexColor("#E8EEF4")))
            if st:
                style.append(("FONTNAME", (0, r), (-1, r), "Helvetica-Bold"))
        for j, (_, t) in enumerate(bloc.colonnes):
            if t in (MONTANT, POURCENT):
                style.append(("ALIGN", (j, 1), (j, -1), "RIGHT"))
        table = LongTable(donnees, colWidths=largeurs, repeatRows=1)
        table.setStyle(TableStyle(style))
        histoire.append(table)
        histoire.append(Spacer(1, 6 * mm))

    edite = datetime.now().strftime("%d/%m/%Y %H:%M")

    def _pied(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.grey)
        canvas.drawString(12 * mm, 8 * mm, _pdf_texte(
            f"{doc.titre} — édité le {edite} — compta-ohada"))
        canvas.drawRightString(taille[0] - 12 * mm, 8 * mm, f"Page {document.page}")
        canvas.restoreState()

    SimpleDocTemplate(str(chemin), pagesize=taille, leftMargin=12 * mm,
                      rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=14 * mm,
                      title=_pdf_texte(doc.titre), author="compta-ohada"
                      ).build(histoire, onFirstPage=_pied, onLaterPages=_pied)


# ---------------------------------------------------------------------------
# Mixin
# ---------------------------------------------------------------------------

class ExportsMixin:
    """``exporter_etat`` de ``Comptabilite``."""

    ETATS_EXPORTABLES = (
        "balance", "grand_livre", "journal", "balance_tiers", "balance_agee",
        "declaration_tva", "bilan", "compte_de_resultat", "tableau_flux_tresorerie",
        "tableau_immobilisations", "plan_amortissement", "etats_smt",
        "livre_recettes_depenses", "resultat_analytique",
    )

    def document_etat(self, nom: str, **params) -> Document:
        """Construit la forme tabulaire d'un état (utile pour un écran)."""
        if nom == "balance":
            return doc_balance(self.balance(params["debut"], params["fin"],
                                            **_sauf(params, "debut", "fin")))
        if nom == "grand_livre":
            return doc_grand_livre(self.grand_livre(params["debut"], params["fin"],
                                                    **_sauf(params, "debut", "fin")),
                                   params["debut"], params["fin"])
        if nom == "journal":
            return doc_journal(self.journal(params["code"], params["debut"],
                                            params["fin"]))
        if nom == "balance_tiers":
            return doc_balance_tiers(self.balance_tiers(params["prefixe"], params["fin"]),
                                     params["prefixe"], params["fin"])
        if nom == "balance_agee":
            return doc_balance_agee(self.balance_agee(params["prefixe"], params["jour"]),
                                    params["prefixe"], params["jour"])
        if nom == "declaration_tva":
            return doc_declaration_tva(self.declaration_tva(params["debut"],
                                                            params["fin"]))
        if nom == "bilan":
            return doc_bilan(self.bilan(params["exercice"],
                                        date_arrete=params.get("date_arrete")))
        if nom == "compte_de_resultat":
            return doc_compte_de_resultat(self.compte_de_resultat(params["exercice"]))
        if nom == "tableau_flux_tresorerie":
            return doc_tft(self.tableau_flux_tresorerie(params["exercice"]))
        if nom == "tableau_immobilisations":
            return doc_tableau_immobilisations(
                self.tableau_immobilisations(params["exercice"]))
        if nom == "plan_amortissement":
            return doc_plan_amortissement(self.plan_amortissement(params["code"]),
                                          params["code"])
        if nom == "etats_smt":
            return doc_etats_smt(self.etats_smt(params["exercice"],
                                                activite=params.get("activite")))
        if nom == "livre_recettes_depenses":
            return doc_livre_recettes_depenses(
                self.livre_recettes_depenses(params["debut"], params["fin"]))
        if nom == "resultat_analytique":
            return doc_resultat_analytique(self.resultat_analytique(
                params["debut"], params["fin"], axe=params.get("axe")))
        raise ParametreInvalide(
            f"État inconnu : {nom!r} (exportables : {', '.join(self.ETATS_EXPORTABLES)}).")

    def _entete_export(self) -> list[str]:
        raison = self.parametre("raison_sociale") or self.tenant_id
        lignes = [raison]
        details = []
        for cle, lib in (("adresse", None), ("nif", "NIF"), ("rccm", "RCCM")):
            v = self.parametre(cle)
            if v:
                details.append(f"{lib} {v}" if lib else v)
        if details:
            lignes.append(" — ".join(details))
        lignes.append(f"Montants en {self.parametre('devise') or 'XOF'}")
        return lignes

    def exporter_etat(self, nom: str, chemin, *, format: str | None = None,
                      **params) -> Path:
        """Écrit l'état ``nom`` dans ``chemin`` (.xlsx ou .pdf) et renvoie le
        chemin. Les paramètres sont ceux de la méthode d'état correspondante
        (``debut``/``fin``, ``exercice``, ``prefixe``, ``code``…)."""
        chemin = Path(chemin)
        format = (format or chemin.suffix.lstrip(".")).lower()
        if format not in ("xlsx", "pdf"):
            raise ParametreInvalide("Format d'export : xlsx ou pdf.")
        doc = self.document_etat(nom, **params)
        devise = self.parametre("devise") or "XOF"
        decimales = 0 if devise in ("XOF", "XAF", "GNF") else 2
        chemin.parent.mkdir(parents=True, exist_ok=True)
        (rendre_xlsx if format == "xlsx" else rendre_pdf)(
            doc, chemin, self._entete_export(), decimales)
        return chemin


def _sauf(params: dict, *cles: str) -> dict:
    return {k: v for k, v in params.items() if k not in cles}
