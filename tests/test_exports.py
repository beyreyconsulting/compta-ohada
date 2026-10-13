"""Exports Excel et PDF de tous les états."""

import pytest

pytest.importorskip("openpyxl")
pytest.importorskip("reportlab")
pypdf = pytest.importorskip("pypdf")

from openpyxl import load_workbook  # noqa: E402

from compta_ohada import ParametreInvalide  # noqa: E402
from conftest import D, J  # noqa: E402
from test_etats_et_cloture import activite_2026  # noqa: E402

PARAMS = {
    "balance": {"debut": "2026-01-01", "fin": "2026-12-31"},
    "grand_livre": {"debut": "2026-01-01", "fin": "2026-12-31", "prefixe": "52"},
    "journal": {"code": "VT", "debut": "2026-01-01", "fin": "2026-12-31"},
    "balance_tiers": {"prefixe": "411", "fin": "2026-12-31"},
    "balance_agee": {"prefixe": "411", "jour": "2026-12-31"},
    "declaration_tva": {"debut": "2026-01-01", "fin": "2026-12-31"},
    "bilan": {"exercice": "2026"},
    "compte_de_resultat": {"exercice": "2026"},
    "tableau_flux_tresorerie": {"exercice": "2026"},
    "tableau_immobilisations": {"exercice": "2026"},
    "plan_amortissement": {"code": "INFO"},
    "etats_smt": {"exercice": "2026", "activite": "negoce"},
    "livre_recettes_depenses": {"debut": "2026-01-01", "fin": "2026-12-31"},
    "resultat_analytique": {"debut": "2026-01-01", "fin": "2026-12-31"},
}


@pytest.fixture
def complet(compta):
    compta.definir_parametre("raison_sociale", "Quincaillerie Alpha SARL")
    compta.definir_parametre("nif", "084123456M")
    compta.definir_parametre("adresse", "ACI 2000, Bamako")
    activite_2026(compta)
    compta.creer_immobilisation("INFO", "Serveur — caisse", "2442",
                                valeur_origine="600000",
                                date_acquisition="2026-03-01", duree_mois=36,
                                contrepartie="5211")
    compta.comptabiliser_dotations("2026-12-31")
    return compta


def test_tous_les_etats_sont_exportables(complet):
    assert set(PARAMS) == set(complet.ETATS_EXPORTABLES)


@pytest.mark.parametrize("nom", sorted(PARAMS))
def test_export_xlsx(complet, tmp_path, nom):
    chemin = complet.exporter_etat(nom, tmp_path / f"{nom}.xlsx", **PARAMS[nom])
    ws = load_workbook(chemin).active
    valeurs = [c for row in ws.iter_rows(values_only=True) for c in row if c is not None]
    assert valeurs[0] == "Quincaillerie Alpha SARL"
    assert "NIF 084123456M" in valeurs[1]
    # Les montants restent des NOMBRES (une valeur entière revient en int).
    assert any(isinstance(v, (int, float)) and not isinstance(v, bool)
               for v in valeurs[2:])


@pytest.mark.parametrize("nom", sorted(PARAMS))
def test_export_pdf(complet, tmp_path, nom):
    chemin = complet.exporter_etat(nom, tmp_path / f"{nom}.pdf", **PARAMS[nom])
    texte = "".join(p.extract_text() for p in pypdf.PdfReader(chemin).pages)
    assert "Quincaillerie Alpha SARL" in texte
    assert "Page 1" in texte


def test_valeurs_numeriques_dans_excel(complet, tmp_path):
    chemin = complet.exporter_etat("bilan", tmp_path / "b.xlsx", exercice="2026")
    ws = load_workbook(chemin).active
    lignes = {r[0]: r for r in ws.iter_rows(values_only=True) if r[0]}
    bilan = complet.bilan("2026")
    assert lignes["BZ"][4] == float(bilan["total_actif"])
    assert lignes["DZ"][2] == float(bilan["total_passif"])


def test_montants_lisibles_dans_le_pdf(complet, tmp_path):
    chemin = complet.exporter_etat("balance", tmp_path / "b.pdf",
                                   debut="2026-01-01", fin="2026-12-31")
    texte = "".join(p.extract_text() for p in pypdf.PdfReader(chemin).pages)
    assert "5 000 000" in texte                 # capital, séparateur de milliers


def test_grand_livre_volumineux_sur_plusieurs_pages(compta, tmp_path):
    from compta_ohada import Ligne
    for i in range(150):
        compta.saisir("CA", J(2026, 1, 1 + i % 28), f"Vente comptant {i}",
                      [Ligne("5711", debit=1000 + i), Ligne("7011", credit=1000 + i)])
    chemin = compta.exporter_etat("grand_livre", tmp_path / "gl.pdf",
                                  debut="2026-01-01", fin="2026-12-31")
    assert len(pypdf.PdfReader(chemin).pages) > 2


def test_export_erreurs(complet, tmp_path):
    with pytest.raises(ParametreInvalide):
        complet.exporter_etat("bilan", tmp_path / "b.docx", exercice="2026")
    with pytest.raises(ParametreInvalide):
        complet.exporter_etat("inconnu", tmp_path / "x.pdf")
